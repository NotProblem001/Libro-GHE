"""Motor de planificación de turnos ProContact.

No es necesario abrir ni modificar este archivo: se usa desde el libro "planificacion_turnos.ipynb".
Todas las reglas que se pueden cambiar están en la plantilla Excel data/PLANTILLA_MES.xlsx.

Reglas implementadas (documento "Horarios y Duración de Turnos.docx"):
- 42 h semanales de lunes a sábado; domingo no se trabaja.
- Colación 60 min L-V; sábado 30 min (10:30 u 11:00) sin pausas.
- Pausas de 15 min tras 2,0-3,0 h conectado; nunca al entrar ni justo antes de salir.
- Máximo 30% de la plataforma en descanso a la vez; 1 en colación por servicio por cada 4 en turno.
- Coberturas mínimas globales (08:00, 09:00, cierre 20:00, sábados).
- Secuencia de sábados: 2 libres, 1 trabajado.
"""
import json
import math
import re
import warnings
from collections import Counter
from datetime import time, timedelta, datetime
from functools import lru_cache
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

DIAS_ES = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
TURNOS_REQUERIDOS = ["Mañana C/sábado", "Tarde C/sábado", "Mañana S/sábado", "Tarde S/sábado", "Intermedio S/sábado"]
SERVICIOS = ["IMAGENES", "EXAMEN Y PROCEDIMIENTOS"]
TIPOS_AUSENCIA = ["LICENCIA", "VACACIONES", "PERMISO"]


class ErrorUsuario(Exception):
    """Problema que el usuario puede corregir en sus archivos (se muestra sin detalles técnicos)."""


# ═══════════════════════════ Valores por defecto de la plantilla ═══════════════════════════
# (clave interna, texto en la hoja PARAMETROS, valor, explicación)
PARAMETROS = [
    ("horas_semanales_contrato", "Horas semanales por contrato", 42, "Jornada legal semanal. Cada turno de la hoja TURNOS debe sumar esto."),
    ("max_horas_diarias", "Máximo de horas por día", 10, "Tope legal de jornada diaria."),
    ("descanso_min_entre_jornadas_h", "Descanso mínimo entre jornadas (horas)", 12, "Horas entre la salida de un día y la entrada del siguiente."),
    ("colacion_min_lv", "Colación lunes a viernes (minutos)", 60, ""),
    ("colacion_min_sab", "Colación sábado (minutos)", 30, ""),
    ("colacion_sab_inicios", "Horarios de colación del sábado", "10:30, 11:00", "Horas de salida a colación permitidas el sábado, separadas por coma."),
    ("pausa_min", "Duración de cada pausa (minutos)", 15, ""),
    ("conectado_min_h", "Pausa: mínimo de horas conectado antes", 2, "La pausa se toma entre 2,0 y 3,0 h después de entrar / volver."),
    ("conectado_max_h", "Pausa: máximo de horas conectado antes", 3, ""),
    ("pausa_margen_salida_min", "Pausa: minutos mínimos antes de la salida", 30, "No se puede tomar la pausa justo antes de salir."),
    ("max_pct_en_descanso", "Máximo % de la plataforma en descanso a la vez", 30, "Incluye pausas y colaciones."),
    ("ejecutivos_por_cupo_colacion", "Ejecutivos en turno por cada cupo de colación (por servicio)", 4, "1 en colación por cada 4 en turno del mismo servicio."),
    ("ciclo_sabados", "Sábados: trabaja 1 de cada", 3, "3 = 2 sábados libres y 1 trabajado."),
    ("cupo_tarde", "Ejecutivos en turno Tarde cada semana", 3, "Asegura la cobertura del cierre (20:00)."),
    ("cupo_manana", "Ejecutivos en turno Mañana cada semana", 6, "Asegura la cobertura de apertura (08:00)."),
]
TURNOS_DEFECTO = {  # [inicio, término] de Lunes a Sábado; None = no trabaja
    "Mañana C/sábado": [["08:00", "16:30"]] * 3 + [["08:00", "16:00"]] * 2 + [["08:00", "14:00"]],
    "Tarde C/sábado": [["11:30", "20:00"]] * 3 + [["12:00", "20:00"]] * 2 + [["08:00", "14:00"]],
    "Mañana S/sábado": [["08:00", "17:30"]] * 4 + [["08:00", "17:00"]] + [None],
    "Tarde S/sábado": [["10:30", "20:00"]] * 4 + [["11:00", "20:00"]] + [None],
    "Intermedio S/sábado": [["09:00", "19:00"]] * 2 + [["09:00", "18:00"]] * 3 + [None],
}
COBERTURA_DEFECTO = [("Lunes a Viernes", "08:00", 6), ("Lunes a Viernes", "09:00", 3), ("Lunes a Viernes", "19:45", 3),
                     ("Sábado", "08:00", 4), ("Sábado", "13:45", 4)]
FERIADOS_DEFECTO = [  # Chile (verificar cada año con el calendario oficial)
    ("2026-01-01", "Año Nuevo"), ("2026-04-03", "Viernes Santo"), ("2026-04-04", "Sábado Santo"),
    ("2026-05-01", "Día del Trabajo"), ("2026-05-21", "Glorias Navales"), ("2026-06-21", "Pueblos Indígenas"),
    ("2026-06-29", "San Pedro y San Pablo"), ("2026-07-16", "Virgen del Carmen"), ("2026-08-15", "Asunción de la Virgen"),
    ("2026-09-18", "Independencia Nacional"), ("2026-09-19", "Glorias del Ejército"), ("2026-10-12", "Encuentro de Dos Mundos"),
    ("2026-10-31", "Iglesias Evangélicas"), ("2026-11-01", "Todos los Santos"), ("2026-12-08", "Inmaculada Concepción"),
    ("2026-12-25", "Navidad"),
]

AZUL, GRIS, BLANCO = "1F3864", "D9D9D9", "FFFFFF"
COLORES_ESTADO = {"LIBRE": "E2EFDA", "LICENCIA": "FCE4D6", "VACACIONES": "DDEBF7", "PERMISO": "EDE7F6", "FERIADO": "FFF2CC"}
BORDE = Border(*(Side(style="thin", color="A6A6A6"),) * 4)
CENTRO = Alignment(horizontal="center", vertical="center")


# ═══════════════════════════ Utilidades ═══════════════════════════
def a_min(hhmm):
    h, m = hhmm.split(":")
    return int(h) * 60 + int(m)


def a_hhmm(minutos):
    return f"{minutos // 60:02d}:{minutos % 60:02d}"


def a_time(hhmm):
    return time(*map(int, hhmm.split(":"))) if isinstance(hhmm, str) else None


def semana_abs(fecha):
    """Nº absoluto de semana (lunes a domingo): la secuencia de sábados continúa entre meses."""
    return (pd.Timestamp(fecha) - pd.Timestamp("2000-01-03")).days // 7


def calcular_dv(cuerpo):
    """Dígito verificador de RUT (módulo 11)."""
    suma, factor = 0, 2
    for d in reversed(str(cuerpo)):
        suma += int(d) * factor
        factor = 2 if factor == 7 else factor + 1
    resto = 11 - suma % 11
    return {11: "0", 10: "K"}.get(resto, str(resto))


def normalizar_rut(valor):
    """'12.345.678-5' | 12345678 | '12345678-5' -> '12345678-5'. Sin dígito verificador, se calcula."""
    if valor is None or pd.isna(valor) or str(valor).strip() == "":
        raise ValueError("RUT vacío")
    if isinstance(valor, (int, float, np.integer, np.floating)):
        cuerpo, dv = str(int(valor)), ""
    else:
        txt = str(valor).upper().replace(".", "").replace(" ", "").replace("\xa0", "")
        cuerpo, _, dv = txt.partition("-")
        if not dv and cuerpo.endswith("K"):
            cuerpo, dv = cuerpo[:-1], "K"
    if not cuerpo.isdigit():
        raise ValueError(f"RUT con formato inválido: {valor}")
    esperado = calcular_dv(int(cuerpo))
    if dv and dv != esperado:
        raise ValueError(f"RUT {valor}: el dígito verificador no corresponde (debería ser {esperado})")
    return f"{int(cuerpo)}-{esperado}"


def normalizar_hora(valor):
    """'HH:MM' o None. Acepta horas de Excel, '8:00', '08:00:00', '8.30'. 00:00 = sin turno."""
    if valor is None or (not isinstance(valor, (time, datetime, timedelta)) and pd.isna(valor)):
        return None
    if isinstance(valor, datetime):
        valor = valor.time()
    if isinstance(valor, time):
        minutos = valor.hour * 60 + valor.minute
    elif isinstance(valor, timedelta):
        minutos = int(valor.total_seconds() // 60)
    elif isinstance(valor, (int, float, np.integer, np.floating)):
        minutos = round(float(valor) * 1440)
    else:
        m = re.fullmatch(r"\s*(\d{1,2})[:.](\d{2})(?::\d{2})?\s*", str(valor))
        if not m:
            return None
        minutos = int(m[1]) * 60 + int(m[2])
    return a_hhmm(minutos) if 0 < minutos < 1440 else None


def limpiar_texto(s):
    return " ".join(str(s).replace("\xa0", " ").split()).upper()


def neto_semanal(dias, reglas):
    return sum(a_min(d[1]) - a_min(d[0]) - (reglas["colacion_min_sab"] if k == 5 else reglas["colacion_min_lv"])
               for k, d in enumerate(dias) if d) / 60


# ═══════════════════════════ Plantilla Excel ═══════════════════════════
def _encabezado(ws, columnas, anchos):
    ws.append(columnas)
    for c, ancho in enumerate(anchos, start=1):
        cel = ws.cell(1, c)
        cel.fill, cel.font, cel.border = PatternFill("solid", fgColor=AZUL), Font(bold=True, color=BLANCO), BORDE
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = ancho
    ws.freeze_panes = "A2"


def _lista(ws, opciones, rango):
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=True,
                        error="Elija un valor de la lista.", showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(rango)


def crear_plantilla(ruta, ejecutivos=None):
    """Crea la plantilla de configuración del mes con los valores del documento de reglas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INSTRUCCIONES"
    texto = [
        "CÓMO PREPARAR EL MES (solo modifique lo que cambió respecto del mes anterior)",
        "",
        "1. EJECUTIVOS: una fila por persona. Si alguien se va, cambie ESTADO a Inactivo. Si llega alguien, agréguelo al final.",
        "   El RUT puede escribirse con o sin puntos y guión. SERVICIO y ESTADO se eligen de la lista.",
        "2. AUSENCIAS: licencias, vacaciones o permisos del mes (una fila por período, fechas DESDE y HASTA).",
        "3. FERIADOS: revise que estén los feriados del mes a planificar.",
        "4. TURNOS, PARAMETROS y COBERTURA: cambie solo si cambian las reglas de la empresa.",
        "   No cambie los nombres de los turnos. Cada turno debe sumar las horas semanales del contrato.",
        "5. Guarde y CIERRE este archivo. Vuelva al libro y presione el botón ▶ (play).",
        "",
        "Si algo está mal, el libro le dirá exactamente qué hoja y qué fila corregir.",
        "Si la plantilla se daña: bórrela, presione ▶ y se creará una nueva con los valores originales.",
    ]
    for linea in texto:
        ws.append([linea])
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("EJECUTIVOS")
    _encabezado(ws, ["RUT", "NOMBRE", "ESTADO", "SERVICIO", "HORAS_CONTRATO"], [14, 42, 12, 28, 16])
    for e in (ejecutivos.itertuples(index=False) if ejecutivos is not None else []):
        ws.append([e.RUT, e.NOMBRE, e.ESTADO, e.SERVICIO, int(e.HORAS_CONTRATO)])
    _lista(ws, ["Activo", "Inactivo"], "C2:C500")
    _lista(ws, SERVICIOS, "D2:D500")

    ws = wb.create_sheet("AUSENCIAS")
    _encabezado(ws, ["RUT", "NOMBRE (referencia)", "TIPO", "DESDE", "HASTA"], [14, 42, 14, 12, 12])
    _lista(ws, TIPOS_AUSENCIA, "C2:C500")
    for col in "DE":
        for fila in range(2, 501):
            ws[f"{col}{fila}"].number_format = "dd-mm-yyyy"

    ws = wb.create_sheet("FERIADOS")
    _encabezado(ws, ["FECHA", "DESCRIPCIÓN"], [14, 30])
    for f, d in FERIADOS_DEFECTO:
        ws.append([pd.Timestamp(f).date(), d])
        ws.cell(ws.max_row, 1).number_format = "dd-mm-yyyy"

    ws = wb.create_sheet("TURNOS")
    _encabezado(ws, ["TURNO"] + [f"{d} {x}" for d in DIAS_ES[:6] for x in ("inicio", "término")] + ["HORAS SEMANA"],
                [22] + [9] * 12 + [10])
    for i, (nombre, dias) in enumerate(TURNOS_DEFECTO.items(), start=2):
        fila = [nombre]
        for d in dias:
            fila += [a_time(d[0]), a_time(d[1])] if d else [None, None]
        # Horas netas de la semana (descuenta colación) para que el usuario vea si suma 42.
        partes = [f"IF(COUNT({get_column_letter(2 + 2 * k)}{i}:{get_column_letter(3 + 2 * k)}{i})=2,"
                  f"({get_column_letter(3 + 2 * k)}{i}-{get_column_letter(2 + 2 * k)}{i})*24-{'0.5' if k == 5 else '1'},0)"
                  for k in range(6)]
        ws.append(fila + ["=" + "+".join(partes)])
        for c in range(2, 14):
            ws.cell(i, c).number_format = "hh:mm"
            ws.cell(i, c).alignment = CENTRO
        ws.cell(i, 14).number_format = "0.0"

    ws = wb.create_sheet("PARAMETROS")
    _encabezado(ws, ["PARÁMETRO", "VALOR", "EXPLICACIÓN"], [55, 14, 80])
    for _, etiqueta, valor, desc in PARAMETROS:
        ws.append([etiqueta, valor, desc or None])

    ws = wb.create_sheet("COBERTURA")
    _encabezado(ws, ["DÍAS", "HORA", "MÍNIMO DE EJECUTIVOS DISPONIBLES"], [18, 10, 34])
    for dias, hora, minimo in COBERTURA_DEFECTO:
        ws.append([dias, a_time(hora), minimo])
        ws.cell(ws.max_row, 2).number_format = "hh:mm"
    _lista(ws, ["Lunes a Viernes", "Sábado"], "A2:A100")

    wb.active = 1
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def leer_plantilla(ruta):
    """Lee y valida la plantilla. Junta TODOS los problemas y los informa de una vez (hoja y fila)."""
    try:
        hojas = pd.read_excel(ruta, sheet_name=None)
    except PermissionError:
        raise ErrorUsuario(f"La plantilla '{ruta.name}' está abierta en Excel. Ciérrela y presione ▶ de nuevo.") from None
    faltan = {"EJECUTIVOS", "AUSENCIAS", "FERIADOS", "TURNOS", "PARAMETROS", "COBERTURA"} - set(hojas)
    if faltan:
        raise ErrorUsuario(f"A la plantilla le faltan las hojas: {', '.join(sorted(faltan))}.\n"
                           "Solución: borre data/PLANTILLA_MES.xlsx y presione ▶ para crear una nueva.")
    for df in hojas.values():
        df.columns = [str(c).strip().upper() for c in df.columns]
    errores, avisos = [], []

    # PARAMETROS
    valores = {limpiar_texto(r.iloc[0]): r.iloc[1] for _, r in hojas["PARAMETROS"].iterrows() if pd.notna(r.iloc[0])}
    reglas = {}
    for clave, etiqueta, defecto, _ in PARAMETROS:
        v = valores.get(limpiar_texto(etiqueta))
        if v is None or pd.isna(v):
            avisos.append(f"PARAMETROS: falta '{etiqueta}', se usa {defecto}.")
            v = defecto
        if clave == "colacion_sab_inicios":
            horas = [normalizar_hora(x) for x in str(v).split(",")]
            if None in horas:
                errores.append(f"PARAMETROS · '{etiqueta}': escriba horas separadas por coma, ej. 10:30, 11:00 (dice '{v}').")
            reglas[clave] = [h for h in horas if h]
            continue
        try:
            reglas[clave] = float(v)
        except (TypeError, ValueError):
            errores.append(f"PARAMETROS · '{etiqueta}': debe ser un número (dice '{v}').")
            reglas[clave] = float(defecto)
    if reglas["max_pct_en_descanso"] > 1:  # se acepta 30 o 30%
        reglas["max_pct_en_descanso"] /= 100
    for clave in ("ciclo_sabados", "cupo_tarde", "cupo_manana", "ejecutivos_por_cupo_colacion", "pausa_min",
                  "colacion_min_lv", "colacion_min_sab", "pausa_margen_salida_min"):
        reglas[clave] = int(reglas[clave])
    reglas["cupos_semana"] = {"Tarde": reglas.pop("cupo_tarde"), "Mañana": reglas.pop("cupo_manana")}

    # TURNOS
    reglas["turnos"] = {}
    for fila, r in zip(hojas["TURNOS"].index + 2, hojas["TURNOS"].itertuples(index=False)):
        if pd.isna(r[0]):
            continue
        nombre = " ".join(str(r[0]).split())
        dias = []
        for k in range(6):
            ini, fin = normalizar_hora(r[1 + 2 * k]), normalizar_hora(r[2 + 2 * k])
            if bool(ini) != bool(fin) or (ini and a_min(fin) <= a_min(ini)):
                errores.append(f"TURNOS fila {fila} · {nombre} · {DIAS_ES[k]}: inicio y término no son válidos.")
                ini = None
            dias.append([ini, fin] if ini else None)
        reglas["turnos"][nombre] = dias
    for nombre in TURNOS_REQUERIDOS:
        dias = reglas["turnos"].get(nombre)
        if dias is None:
            errores.append(f"TURNOS: falta el turno '{nombre}' (no cambie los nombres de los turnos).")
            continue
        if ("C/sábado" in nombre) != (dias[5] is not None):
            errores.append(f"TURNOS · {nombre}: {'debe' if 'C/sábado' in nombre else 'no debe'} tener horario el sábado.")
        neto = neto_semanal(dias, reglas)
        if abs(neto - reglas["horas_semanales_contrato"]) > 1e-9:
            errores.append(f"TURNOS · {nombre}: suma {neto:g} h netas a la semana y deben ser "
                           f"{reglas['horas_semanales_contrato']:g} (columna HORAS SEMANA).")

    # COBERTURA
    reglas["cobertura_min"] = {"LV": {}, "SAB": {}}
    for fila, r in zip(hojas["COBERTURA"].index + 2, hojas["COBERTURA"].itertuples(index=False)):
        if pd.isna(r[0]):
            continue
        hora = normalizar_hora(r[1])
        if hora is None or pd.isna(r[2]):
            errores.append(f"COBERTURA fila {fila}: falta la hora o el mínimo.")
            continue
        reglas["cobertura_min"]["SAB" if "SÁB" in limpiar_texto(r[0]) or "SAB" in limpiar_texto(r[0]) else "LV"][hora] = int(r[2])

    # FERIADOS
    feriados = pd.to_datetime(hojas["FERIADOS"]["FECHA"], errors="coerce", dayfirst=True)
    if feriados.isna().any():
        errores.append("FERIADOS: hay fechas inválidas en las filas " + ", ".join(str(i + 2) for i in feriados[feriados.isna()].index) + ".")

    # EJECUTIVOS
    ej = hojas["EJECUTIVOS"].dropna(how="all").copy()
    for col in ("RUT", "NOMBRE"):
        if col not in ej:
            errores.append(f"EJECUTIVOS: falta la columna {col}.")
    if not errores or "RUT" in ej:
        ruts = []
        for fila, v in zip(ej.index + 2, ej.get("RUT", [])):
            try:
                ruts.append(normalizar_rut(v))
            except ValueError as e:
                errores.append(f"EJECUTIVOS fila {fila}: {e}.")
                ruts.append(None)
        ej["RUT"] = ruts
        ej["NOMBRE"] = ej["NOMBRE"].fillna("").map(limpiar_texto)
        ej["ESTADO"] = ej.get("ESTADO", pd.Series("Activo", index=ej.index)).fillna("Activo").astype(str).str.strip().str.capitalize()
        ej["SERVICIO"] = ej.get("SERVICIO", pd.Series(index=ej.index, dtype=object)).fillna("SIN ASIGNAR").map(limpiar_texto)
        ej["HORAS_CONTRATO"] = pd.to_numeric(ej.get("HORAS_CONTRATO"), errors="coerce")
        ej["HORAS_CONTRATO"] = ej["HORAS_CONTRATO"].fillna(reglas["horas_semanales_contrato"])
        dup = ej[ej["RUT"].notna() & ej["RUT"].duplicated(keep=False)]
        for rut, g in dup.groupby("RUT"):
            errores.append(f"EJECUTIVOS: el RUT {rut} está repetido en las filas {', '.join(str(i + 2) for i in g.index)}.")
        for fila, r in zip(ej.index + 2, ej.itertuples(index=False)):
            if r.ESTADO == "Activo" and r.HORAS_CONTRATO != reglas["horas_semanales_contrato"]:
                avisos.append(f"EJECUTIVOS fila {fila}: {r.NOMBRE} tiene contrato de {r.HORAS_CONTRATO:g} h; "
                              f"los turnos son de {reglas['horas_semanales_contrato']:g} h (aparecerá en ALERTAS).")
    activos = ej[ej["ESTADO"] == "Activo"].sort_values("NOMBRE").reset_index(drop=True) if not errores else ej
    if not errores and activos.empty:
        errores.append("EJECUTIVOS: no hay ninguna persona con ESTADO = Activo.")

    # AUSENCIAS
    ausencias = {}
    au = hojas["AUSENCIAS"].dropna(subset=[c for c in ("RUT", "TIPO", "DESDE", "HASTA") if c in hojas["AUSENCIAS"]], how="all")
    ruts_activos = set(activos["RUT"]) if not errores else set()
    for fila, r in zip(au.index + 2, au.to_dict("records")):
        try:
            rut = normalizar_rut(r.get("RUT"))
        except ValueError as e:
            errores.append(f"AUSENCIAS fila {fila}: {e}.")
            continue
        desde, hasta = pd.to_datetime(r.get("DESDE"), errors="coerce", dayfirst=True), pd.to_datetime(r.get("HASTA"), errors="coerce", dayfirst=True)
        if pd.isna(desde) or pd.isna(hasta) or hasta < desde:
            errores.append(f"AUSENCIAS fila {fila}: revise las fechas DESDE y HASTA.")
            continue
        if ruts_activos and rut not in ruts_activos:
            avisos.append(f"AUSENCIAS fila {fila}: el RUT {rut} no está activo en EJECUTIVOS; se ignora.")
            continue
        tipo = limpiar_texto(r.get("TIPO")) if pd.notna(r.get("TIPO")) else "LICENCIA"
        for d in pd.date_range(desde, hasta):
            ausencias[(rut, d)] = tipo

    if errores:
        raise ErrorUsuario("Corrija lo siguiente en data/PLANTILLA_MES.xlsx, guarde, cierre el archivo y presione ▶ de nuevo:\n"
                           + "\n".join(f"   • {e}" for e in errores))
    return activos[["RUT", "NOMBRE", "ESTADO", "SERVICIO", "HORAS_CONTRATO"]], ausencias, reglas, feriados.tolist(), avisos


# ═══════════════════════════ Horario del mes anterior ═══════════════════════════
def cargar_historico(rutas):
    """Turnos previos (RUT, FECHA, H_INICIO, H_TERMINO, TURNO) de todas las hojas con encabezado RUT/FECHA/H_INICIO.
    Sirve el Excel histórico de ProContact y las planillas generadas por este libro."""
    columnas = ["RUT", "FECHA", "H_INICIO", "H_TERMINO", "TURNO"]
    partes = []
    for orden, ruta in enumerate(rutas):
        try:
            hojas = pd.read_excel(ruta, sheet_name=None, header=None)
        except PermissionError:
            raise ErrorUsuario(f"'{ruta.name}' está abierto en Excel. Ciérrelo y presione ▶ de nuevo.") from None
        except Exception as e:
            raise ErrorUsuario(f"No se pudo leer '{ruta.name}' ({e}). ¿Es un archivo Excel válido?") from None
        for nombre, crudo in hojas.items():
            fila = next((i for i in range(min(30, len(crudo)))
                         if {"RUT", "FECHA", "H_INICIO"} <= {str(v).strip().upper() for v in crudo.iloc[i]}), None)
            if fila is None:
                continue
            df = crudo.iloc[fila + 1:].copy()
            df.columns = [str(c).strip().upper() for c in crudo.iloc[fila]]
            df = df.loc[:, ~df.columns.duplicated()]
            partes.append(df.reindex(columns=columnas).assign(ORDEN=orden, HOJA=nombre))
    if not partes:
        return pd.DataFrame(columns=columnas)
    hist = pd.concat(partes, ignore_index=True).dropna(subset=["RUT"])
    ruts = []
    for v in hist["RUT"]:
        try:
            ruts.append(normalizar_rut(v))
        except ValueError:
            ruts.append(None)
    hist["RUT"] = ruts
    hist["FECHA"] = pd.to_datetime(hist["FECHA"], errors="coerce")
    hist = hist.dropna(subset=["RUT", "FECHA"]).copy()
    hist["H_INICIO"] = hist["H_INICIO"].map(normalizar_hora)
    hist["H_TERMINO"] = hist["H_TERMINO"].map(normalizar_hora)
    hist = hist.sort_values(["ORDEN", "FECHA", "HOJA"])  # si una fecha se repite, gana el archivo más reciente
    return hist.drop_duplicates(["RUT", "FECHA"], keep="last")[columnas].reset_index(drop=True)


def familia_turno(turno, inicio_min):
    if isinstance(turno, str) and turno.split()[0] in ("Mañana", "Intermedio", "Tarde"):
        return turno.split()[0]
    return "Tarde" if inicio_min >= a_min("10:30") else "Intermedio" if inicio_min >= a_min("09:00") else "Mañana"


def estado_previo(hist):
    """Por RUT: último sábado trabajado, turnos acumulados por tipo, último turno y fin de la última jornada."""
    previo = {}
    trab = hist[hist["H_INICIO"].notna() & hist["H_TERMINO"].notna()].copy()
    trab["DOW"] = trab["FECHA"].dt.dayofweek
    trab["SEMANA"] = trab["FECHA"].map(semana_abs)
    for rut, g in trab.groupby("RUT"):
        sabados, lv = g[g["DOW"] == 5], g[g["DOW"] < 5]
        familias = [familia_turno(s["TURNO"].iloc[0], a_min(s["H_INICIO"].mode().iloc[0])) for _, s in lv.groupby("SEMANA")]
        ultimo = g.iloc[-1]
        previo[rut] = {
            "semana_ult_sabado": int(sabados["SEMANA"].max()) if not sabados.empty else None,
            "conteo": Counter(familias),
            "ultimo_turno": familias[-1] if familias else None,
            "fin_ultima_jornada": ultimo["FECHA"] + pd.Timedelta(minutes=a_min(ultimo["H_TERMINO"])),
        }
    return previo


# ═══════════════════════════ Motor de asignación ═══════════════════════════
def generar_calendario(inicio, feriados):
    cal = pd.DataFrame({"FECHA": pd.date_range(inicio, inicio + pd.offsets.MonthEnd(0))})
    cal["DOW"] = cal["FECHA"].dt.dayofweek
    cal["DIA"] = cal["DOW"].map(dict(enumerate(DIAS_ES)))
    cal["SEMANA"] = cal["FECHA"].map(semana_abs)
    cal["INICIO_SEMANA"] = cal["FECHA"] - pd.to_timedelta(cal["DOW"], unit="D")
    cal["TIPO_DIA"] = np.select([cal["FECHA"].isin(pd.to_datetime(feriados)), cal["DOW"] == 6, cal["DOW"] == 5],
                                ["FERIADO", "DOMINGO", "SÁBADO"], "HÁBIL")
    return cal


def asignar_fase_sabado(ruts, previo, ciclo, minimo):
    """'Cada 2 sábados libres, trabajar 1': trabaja el sábado de la semana s si (s + fase) % ciclo == 0.
    La fase se hereda del mes anterior; los nuevos van al grupo más chico."""
    fase, nuevos = {}, []
    for rut in ruts:
        s = previo.get(rut, {}).get("semana_ult_sabado")
        if s is None:
            nuevos.append(rut)
        else:
            fase[rut] = (-s) % ciclo
    tam = lambda: Counter({f: 0 for f in range(ciclo)} | Counter(fase.values()))
    for rut in nuevos:
        t = tam()
        fase[rut] = min(t, key=lambda f: (t[f], f))
    while True:  # rebalanceo si un sábado queda bajo la cobertura mínima
        t = tam()
        chico, grande = min(t, key=t.get), max(t, key=t.get)
        if t[chico] >= minimo or t[grande] - t[chico] <= 1:
            break
        rut = sorted((r for r in fase if fase[r] == grande), key=lambda r: (r not in nuevos, r))[0]
        fase[rut] = chico
    return fase


def asignar_turnos_semanales(ejecutivos, calendario, previo, ausencias, reglas):
    """Un turno por ejecutivo y semana: cupos de Tarde y Mañana para quienes menos los han tenido."""
    ciclo, cupos = reglas["ciclo_sabados"], reglas["cupos_semana"]
    ruts = list(ejecutivos["RUT"])
    fase = asignar_fase_sabado(ruts, previo, ciclo, max(reglas["cobertura_min"]["SAB"].values(), default=0))
    conteo = {r: Counter(previo.get(r, {}).get("conteo", {})) for r in ruts}
    ultimo = {r: previo.get(r, {}).get("ultimo_turno") for r in ruts}
    filas = []
    for semana, dias in calendario[calendario["DOW"] < 6].groupby("SEMANA"):
        disponibles = [r for r in ruts if not all((r, f) in ausencias for f in dias["FECHA"])]
        grupo_sab = {r for r in disponibles if (semana + fase[r]) % ciclo == 0}
        tarde = sorted(disponibles, key=lambda r: (conteo[r]["Tarde"], ultimo[r] == "Tarde", r))[:cupos["Tarde"]]
        resto = [r for r in disponibles if r not in tarde]
        manana = [r for r in resto if r in grupo_sab]  # no existe Intermedio con sábado
        manana += sorted((r for r in resto if r not in grupo_sab), key=lambda r: (conteo[r]["Mañana"], r))[:max(0, cupos["Mañana"] - len(manana))]
        for r in ruts:
            familia = "Tarde" if r in tarde else "Mañana" if r in manana else "Intermedio"
            filas.append({"SEMANA": semana, "RUT": r, "TURNO": f"{familia} {'C' if r in grupo_sab else 'S'}/sábado"})
            if r in disponibles:
                conteo[r][familia] += 1
                ultimo[r] = familia
    return pd.DataFrame(filas)


def expandir_plan(ejecutivos, calendario, turnos_semana, ausencias, reglas):
    """Una fila por ejecutivo y día L-S. Prioridad: ausencia > feriado > horario del turno."""
    base = calendario[calendario["DOW"] < 6].merge(turnos_semana, on="SEMANA").merge(ejecutivos, on="RUT")
    def dia(r):
        if (r.RUT, r.FECHA) in ausencias:
            return ausencias[(r.RUT, r.FECHA)], None, None
        if r.TIPO_DIA == "FERIADO":
            return "FERIADO", None, None
        horario = reglas["turnos"][r.TURNO][r.DOW]
        return ("TRABAJA", *horario) if horario else ("LIBRE", None, None)
    base[["ESTADO_DIA", "H_INICIO", "H_TERMINO"]] = pd.DataFrame([dia(r) for r in base.itertuples()], index=base.index)
    return base.sort_values(["NOMBRE", "FECHA"]).reset_index(drop=True)


@lru_cache(maxsize=None)
def opciones_descanso(ini, fin, sabado, reglas_json):
    """Combinaciones válidas (pausa1, salida colación, minutos colación, pausa2) en bloques de 15 min."""
    r = json.loads(reglas_json)
    if sabado:
        dur = r["colacion_min_sab"]
        return [(None, c, dur, None) for c in map(a_min, r["colacion_sab_inicios"]) if ini < c and c + dur <= fin]
    p, col = r["pausa_min"], r["colacion_min_lv"]
    lo, hi = int(r["conectado_min_h"] * 60), int(r["conectado_max_h"] * 60)
    out = []
    for d1 in range(ini + lo, ini + hi + 1, 15):
        for c in range(d1 + p + lo, d1 + p + hi + 1, 15):
            for d2 in range(c + col + lo, c + col + hi + 1, 15):
                if r["pausa_margen_salida_min"] <= fin - (d2 + p) <= hi:
                    out.append((d1, c, col, d2))
    return out


def slots(d1, c, dur, d2, p=15):
    return [s for s in (d1, d2) if s is not None], list(range(c, c + dur, p))


def asignar_descansos(plan, reglas, reinicios=40):
    """Ubica pausas y colación de cada persona donde menos gente esté en descanso (aplana la curva),
    respetando el máximo de 30% y el cupo de colación por servicio."""
    reglas_json = json.dumps(reglas, sort_keys=True)
    plan = plan.copy()
    for col in ["DESCANSO 1", "DESCANSO 2", "SALIDA COLACION", "ENTRADA COLACION"]:
        plan[col] = None
    plan["Minutos Colacion"] = 0
    for fecha, dia in plan[plan["ESTADO_DIA"] == "TRABAJA"].groupby("FECHA"):
        cupo = max(1, math.floor(reglas["max_pct_en_descanso"] * len(dia)))
        cupo_col = {s: max(1, n // reglas["ejecutivos_por_cupo_colacion"]) for s, n in dia["SERVICIO"].value_counts().items()}
        filas = {idx: f for idx, f in dia.sort_values(["H_INICIO", "RUT"]).iterrows()}
        opciones = {}
        for idx, f in filas.items():
            opciones[idx] = opciones_descanso(a_min(f["H_INICIO"]), a_min(f["H_TERMINO"]), f["DOW"] == 5, reglas_json)
            if not opciones[idx]:
                raise ErrorUsuario(f"El turno '{f['TURNO']}' ({f['H_INICIO']} a {f['H_TERMINO']}) no permite ubicar pausas "
                                   "y colación con las reglas de la hoja PARAMETROS. Revise el horario en la hoja TURNOS.")

        def costo(o, serv, ocup, ocup_col):
            pausas, colac = slots(*o)
            c = sum(ocup[s] for s in pausas + colac)
            c += 1000 * sum(ocup[s] + 1 > cupo for s in pausas + colac)
            c += 100 * sum(ocup_col[(serv, s)] + 1 > cupo_col[serv] for s in colac)  # cupo por servicio: regla blanda
            return c

        def resolver(orden):
            ocup, ocup_col, eleccion = Counter(), Counter(), {}
            for _ in range(5):  # 1ª pasada greedy, luego búsqueda local
                for idx in orden:
                    serv = filas[idx]["SERVICIO"]
                    for o, signo in ([(eleccion[idx], -1)] if idx in eleccion else []) + [(None, +1)]:
                        if o is None:
                            o = eleccion[idx] = min(opciones[idx], key=lambda x: costo(x, serv, ocup, ocup_col))
                        pausas, colac = slots(*o)
                        for s in pausas + colac:
                            ocup[s] += signo
                        for s in colac:
                            ocup_col[(serv, s)] += signo
            return sum(max(0, n - cupo) for n in ocup.values()), eleccion

        orden = sorted(filas, key=lambda i: len(opciones[i]))  # turnos más restringidos primero
        mejor = resolver(orden)
        rng = np.random.default_rng(int(pd.Timestamp(fecha).strftime("%Y%m%d")))  # semilla fija: resultado reproducible
        for _ in range(reinicios):
            if mejor[0] == 0:
                break
            mejor = min(mejor, resolver(list(rng.permutation(orden))), key=lambda x: x[0])
        for idx, (d1, c, dur, d2) in mejor[1].items():
            plan.loc[idx, ["DESCANSO 1", "DESCANSO 2", "SALIDA COLACION", "ENTRADA COLACION", "Minutos Colacion"]] = [
                a_hhmm(d1) if d1 is not None else None, a_hhmm(d2) if d2 is not None else None, a_hhmm(c), a_hhmm(c + dur), dur]
    plan["Horas diarias"] = [(a_min(f) - a_min(i)) / 60 if isinstance(i, str) else 0 for i, f in zip(plan["H_INICIO"], plan["H_TERMINO"])]
    plan["Total Horas Trabajadas"] = plan["Horas diarias"] - plan["Minutos Colacion"] / 60
    return plan


# ═══════════════════════════ Validación ═══════════════════════════
def auditar(plan, ejecutivos, previo, reglas):
    """Control de calidad independiente del motor: horas, descanso entre jornadas, pausas y cobertura."""
    alertas = []
    def alerta(tipo, fecha, rut, detalle):
        alertas.append({"TIPO": tipo, "FECHA": fecha, "RUT": rut, "DETALLE": detalle})
    trab = plan[plan["ESTADO_DIA"] == "TRABAJA"].copy()
    contrato = ejecutivos.set_index("RUT")["HORAS_CONTRATO"]

    for _, f in trab[trab["Horas diarias"] > reglas["max_horas_diarias"]].iterrows():
        alerta("EXCESO DE HORAS EN EL DÍA", f["FECHA"], f["RUT"], f"{f['Horas diarias']:.1f} h de jornada")
    semanal = plan.groupby(["RUT", "SEMANA"]).agg(
        horas=("Total Horas Trabajadas", "sum"), dias=("FECHA", "size"), inicio=("INICIO_SEMANA", "first"),
        completa=("ESTADO_DIA", lambda s: s.isin(["TRABAJA", "LIBRE"]).all()))
    for (rut, _), s in semanal.iterrows():
        if s.horas > contrato[rut] + 1e-9:
            alerta("EXCESO DE HORAS EN LA SEMANA", s.inicio, rut, f"{s.horas:.1f} h > contrato {contrato[rut]:g} h")
        elif s.dias == 6 and s.completa and s.horas < contrato[rut] - 1e-9:
            alerta("FALTAN HORAS EN LA SEMANA", s.inicio, rut, f"{s.horas:.1f} h < contrato {contrato[rut]:g} h")

    trab["INI_DT"] = trab["FECHA"] + pd.to_timedelta(trab["H_INICIO"].map(a_min), unit="min")
    trab["FIN_DT"] = trab["FECHA"] + pd.to_timedelta(trab["H_TERMINO"].map(a_min), unit="min")
    trab = trab.sort_values(["RUT", "INI_DT"])
    fin_mes_anterior = pd.to_datetime(trab["RUT"].map({r: p["fin_ultima_jornada"] for r, p in previo.items()}))
    fin_previo = trab.groupby("RUT")["FIN_DT"].shift().fillna(fin_mes_anterior)
    gap_h = (trab["INI_DT"] - fin_previo).dt.total_seconds() / 3600
    for (_, f), g in zip(trab.iterrows(), gap_h):
        if pd.notna(g) and g < reglas["descanso_min_entre_jornadas_h"]:
            alerta("POCO DESCANSO ENTRE JORNADAS", f["FECHA"], f["RUT"], f"solo {g:.1f} h desde la jornada anterior")
        hitos = [a_min(f["H_INICIO"])]
        for ini, dur in [(f["DESCANSO 1"], reglas["pausa_min"]), (f["SALIDA COLACION"], int(f["Minutos Colacion"])),
                         (f["DESCANSO 2"], reglas["pausa_min"])]:
            if isinstance(ini, str):
                hitos += [a_min(ini), a_min(ini) + dur]
        hitos.append(a_min(f["H_TERMINO"]))
        tramos = [hitos[i + 1] - hitos[i] for i in range(0, len(hitos), 2)]
        if max(tramos) > reglas["conectado_max_h"] * 60:
            alerta("MÁS DE 3 H CONECTADO SIN PAUSA", f["FECHA"], f["RUT"], f"tramos de {tramos} min")

    registros = []
    for f in trab.to_dict("records"):
        en_pausa = set()
        for ini, dur in [(f["DESCANSO 1"], reglas["pausa_min"]), (f["DESCANSO 2"], reglas["pausa_min"]),
                         (f["SALIDA COLACION"], int(f["Minutos Colacion"]))]:
            if isinstance(ini, str):
                en_pausa |= set(range(a_min(ini), a_min(ini) + dur, 15))
        for s in range(a_min(f["H_INICIO"]), a_min(f["H_TERMINO"]), 15):
            registros.append((f["FECHA"], s, s not in en_pausa))
    bloques = pd.DataFrame(registros, columns=["FECHA", "BLOQUE", "DISPONIBLE"])
    resumen = bloques.groupby(["FECHA", "BLOQUE"]).agg(EN_TURNO=("DISPONIBLE", "size"), DISPONIBLES=("DISPONIBLE", "sum")).reset_index()
    resumen["EN_DESCANSO"] = resumen["EN_TURNO"] - resumen["DISPONIBLES"]
    en_turno_dia = trab.groupby("FECHA").size()
    for f in resumen.itertuples():
        cupo = max(1, math.floor(reglas["max_pct_en_descanso"] * en_turno_dia[f.FECHA]))
        if f.EN_DESCANSO > cupo:
            alerta("DEMASIADOS EN DESCANSO A LA VEZ", f.FECHA, None, f"{a_hhmm(f.BLOQUE)}: {f.EN_DESCANSO} en descanso (máximo {cupo})")
    disp = resumen.set_index(["FECHA", "BLOQUE"])["DISPONIBLES"]
    feriados = set(plan.loc[plan["ESTADO_DIA"] == "FERIADO", "FECHA"])
    for fecha in sorted(set(plan["FECHA"]) - feriados):
        for hhmm, minimo in reglas["cobertura_min"]["SAB" if fecha.dayofweek == 5 else "LV"].items():
            n = int(disp.get((fecha, a_min(hhmm)), 0))
            if n < minimo:
                alerta("COBERTURA BAJO EL MÍNIMO", fecha, None, f"{hhmm}: {n} disponibles (mínimo {minimo})")

    alertas = pd.DataFrame(alertas, columns=["TIPO", "FECHA", "RUT", "DETALLE"])
    alertas["NOMBRE"] = alertas["RUT"].map(ejecutivos.set_index("RUT")["NOMBRE"])
    return alertas.sort_values(["TIPO", "FECHA"]).reset_index(drop=True)


# ═══════════════════════════ Exportación a Excel ═══════════════════════════
def estilo_tabla(ws, n_filas, n_cols, anchos, fmt_cols=None, congelar="D2"):
    for c in range(1, n_cols + 1):
        cel = ws.cell(1, c)
        cel.fill, cel.font = PatternFill("solid", fgColor=AZUL), Font(bold=True, color=BLANCO)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = anchos[c - 1] if c <= len(anchos) else 12
    for fila in ws.iter_rows(min_row=1, max_row=1 + n_filas, max_col=n_cols):
        for cel in fila:
            cel.border = BORDE
            if cel.row > 1 and cel.column > 1:
                cel.alignment = CENTRO
    for col, fmt in (fmt_cols or {}).items():
        for (cel,) in ws.iter_rows(min_row=2, max_row=1 + n_filas, min_col=col, max_col=col):
            cel.number_format = fmt
    ws.freeze_panes = congelar  # asignar una sola vez: openpyxl acumula selecciones y Excel rechaza el archivo
    ws.auto_filter.ref = f"A1:{get_column_letter(n_cols)}{1 + n_filas}"


def fila_total(ws, fila, n_cols):
    for c in range(1, n_cols + 1):
        ws.cell(fila, c).font, ws.cell(fila, c).fill = Font(bold=True), PatternFill("solid", fgColor=GRIS)


def exportar_excel(ruta, plan, ejecutivos, turnos_semana, calendario, alertas, reglas):
    """4 hojas: PLANIFICACION_MENSUAL (esquema ProContact), RESUMEN_EJECUTIVOS, MATRIZ_COBERTURA, ALERTAS.
    Ojo: nunca escribir "" en una celda (openpyxl genera celdas que Excel rechaza); usar None."""
    wb = openpyxl.Workbook()
    P = "PLANIFICACION_MENSUAL!"

    ws = wb.active
    ws.title = "PLANIFICACION_MENSUAL"
    enc = ["RUT", "Estado", "NOMBRE", "SERVICIO", "FECHA", "DIA", "SEMANA", "TURNO", "ESTADO_DIA", "H_INICIO", "H_TERMINO",
           "DURACION", "Horas diarias", "Minutos Colacion", "hora colacion", "Total Horas Trabajadas", "HORAS TRABAJADAS",
           "DESCANSO 1", "DESCANSO 2", "SALIDA COLACION", "ENTRADA COLACION", "TIEMPO COLACION"]
    ws.append(enc)
    for i, g in enumerate(plan.to_dict("records"), start=2):
        ws.append([
            g["RUT"], g["ESTADO"], g["NOMBRE"], g["SERVICIO"], g["FECHA"].date(), g["DIA"], g["INICIO_SEMANA"].date(),
            g["TURNO"], g["ESTADO_DIA"], a_time(g["H_INICIO"]), a_time(g["H_TERMINO"]),
            f"=IF(COUNT(J{i}:K{i})=2,K{i}-J{i},0)", f"=L{i}*24", int(g["Minutos Colacion"]), f"=N{i}/60", f"=M{i}-O{i}",
            f"=SUMIFS($P:$P,$A:$A,A{i},$G:$G,G{i})",
            a_time(g["DESCANSO 1"]), a_time(g["DESCANSO 2"]), a_time(g["SALIDA COLACION"]), a_time(g["ENTRADA COLACION"]),
            f"=IF(COUNT(T{i}:U{i})=2,ROUND((U{i}-T{i})*1440,0),0)",
        ])
    n_plan = len(plan)
    estilo_tabla(ws, n_plan, len(enc), [12, 9, 38, 26, 12, 11, 12, 20, 12, 9, 10, 10, 9, 10, 9, 11, 11, 10, 10, 11, 11, 10],
                 {c: "hh:mm" for c in (10, 11, 12, 18, 19, 20, 21)} | {5: "yyyy-mm-dd", 7: "yyyy-mm-dd", 13: "0.0", 15: "0.0", 16: "0.0", 17: "0.0"})
    for estado, color in COLORES_ESTADO.items():
        ws.conditional_formatting.add(f"A2:V{n_plan + 1}", FormulaRule(formula=[f'$I2="{estado}"'], fill=PatternFill("solid", fgColor=color)))

    ws = wb.create_sheet("RESUMEN_EJECUTIVOS")
    enc = ["RUT", "NOMBRE", "SERVICIO", "HORAS CONTRATO SEMANAL", "MÁX. HORAS SEMANA", "HORAS PLANIFICADAS MES",
           "HORAS CONTRATO MES (ref.)", "DIFERENCIA", "DÍAS TRABAJADOS", "SÁBADOS TRABAJADOS", "SEMANAS TARDE",
           "DÍAS LIBRE", "LICENCIA", "VACACIONES"]
    ws.append(enc)
    habiles = int((calendario["TIPO_DIA"] == "HÁBIL").sum())
    max_sem = plan.groupby(["RUT", "SEMANA"])["Total Horas Trabajadas"].sum().groupby("RUT").max()
    tardes = turnos_semana[turnos_semana["TURNO"].str.startswith("Tarde")].groupby("RUT").size()
    for i, e in enumerate(ejecutivos.itertuples(index=False), start=2):
        ws.append([
            e.RUT, e.NOMBRE, e.SERVICIO, e.HORAS_CONTRATO, round(float(max_sem.get(e.RUT, 0)), 1),
            f"=SUMIFS({P}$P:$P,{P}$A:$A,A{i})", f"=D{i}/5*{habiles}", f"=F{i}-G{i}",
            f'=COUNTIFS({P}$A:$A,A{i},{P}$I:$I,"TRABAJA")',
            f'=COUNTIFS({P}$A:$A,A{i},{P}$F:$F,"Sábado",{P}$I:$I,"TRABAJA")', int(tardes.get(e.RUT, 0)),
            f'=COUNTIFS({P}$A:$A,A{i},{P}$I:$I,"LIBRE")', f'=COUNTIFS({P}$A:$A,A{i},{P}$I:$I,"LICENCIA")',
            f'=COUNTIFS({P}$A:$A,A{i},{P}$I:$I,"VACACIONES")',
        ])
    n = len(ejecutivos)
    ws.append(["TOTAL", f'=COUNTA(A2:A{n + 1})&" ejecutivos"', None] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{n + 1})" for c in range(4, len(enc) + 1)])
    ws.cell(n + 2, 5).value = f"=MAX(E2:E{n + 1})"
    estilo_tabla(ws, n + 1, len(enc), [12, 38, 26, 12, 11, 13, 13, 11, 10, 10, 10, 9, 9, 11], {c: "0.0" for c in (4, 5, 6, 7, 8)}, congelar="C2")
    fila_total(ws, n + 2, len(enc))
    ws.conditional_formatting.add(f"E2:E{n + 1}", FormulaRule(formula=["E2>D2"], fill=PatternFill("solid", fgColor="F8CBAD")))

    ws = wb.create_sheet("MATRIZ_COBERTURA")
    tramos = [d for dias in reglas["turnos"].values() for d in dias if d]
    horas = list(range(min(a_min(d[0]) for d in tramos) // 60, -(-max(a_min(d[1]) for d in tramos) // 60)))
    ws.append(["FECHA", "DIA", "TIPO_DIA"] + [time(h, 0) for h in horas] + ["TURNOS DEL DÍA"])
    ultima = n_plan + 1
    dias = calendario[calendario["DOW"] < 6].reset_index(drop=True)
    for i, d in enumerate(dias.itertuples(index=False), start=2):
        fila = [d.FECHA.date(), d.DIA, d.TIPO_DIA]
        for k in range(len(horas)):
            col = get_column_letter(4 + k)  # +0,0001 día (~9 s) evita errores de redondeo al comparar horas
            fila.append(f'=COUNTIFS({P}$E$2:$E${ultima},$A{i},{P}$J$2:$J${ultima},"<="&({col}$1+0.0001),'
                        f'{P}$K$2:$K${ultima},">"&({col}$1+0.0001))')
        fila.append(f'=COUNTIFS({P}$E$2:$E${ultima},$A{i},{P}$I$2:$I${ultima},"TRABAJA")')
        ws.append(fila)
    n = len(dias)
    ws.append(["MÍNIMO", None, None] + [f"=MIN({get_column_letter(c)}2:{get_column_letter(c)}{n + 1})" for c in range(4, 5 + len(horas))])
    estilo_tabla(ws, n + 1, 4 + len(horas), [12, 11, 10] + [7] * len(horas) + [10], {1: "yyyy-mm-dd"})
    for c in range(4, 4 + len(horas)):
        ws.cell(1, c).number_format = "hh:mm"
    fila_total(ws, n + 2, 4 + len(horas))
    rojo = PatternFill("solid", fgColor="F8CBAD")
    for tipo, mins in reglas["cobertura_min"].items():
        cond = '$C2="SÁBADO"' if tipo == "SAB" else '$C2="HÁBIL"'
        for hhmm, minimo in mins.items():
            if int(hhmm[:2]) in horas:
                col = get_column_letter(4 + horas.index(int(hhmm[:2])))
                ws.conditional_formatting.add(f"{col}2:{col}{n + 1}", FormulaRule(formula=[f"AND({cond},{col}2<{minimo})"], fill=rojo))

    ws = wb.create_sheet("ALERTAS")
    ws.append(["TIPO", "FECHA", "RUT", "NOMBRE", "DETALLE"])
    for a in alertas.itertuples(index=False):
        ws.append([a.TIPO, pd.Timestamp(a.FECHA).date(), a.RUT if isinstance(a.RUT, str) else None,
                   a.NOMBRE if isinstance(a.NOMBRE, str) else None, a.DETALLE])
    if alertas.empty:
        ws.append(["Sin alertas", None, None, None, "La planificación cumple todas las reglas."])
    estilo_tabla(ws, max(1, len(alertas)), 5, [34, 12, 12, 38, 60], {2: "yyyy-mm-dd"}, congelar="A2")

    try:
        wb.save(ruta)
    except PermissionError:
        raise ErrorUsuario(f"No se pudo guardar '{ruta.name}' porque está abierto en Excel. Ciérrelo y presione ▶ de nuevo.") from None


# ═══════════════════════════ Programa principal ═══════════════════════════
def _mes_desde_texto(mes):
    partes = str(mes).lower().split()
    if len(partes) != 2 or partes[0] not in MESES or not partes[1].isdigit():
        raise ErrorUsuario(f"MES = '{mes}' no se entiende. Escriba \"auto\" o un mes y año, por ejemplo \"Diciembre 2026\".")
    return pd.Timestamp(int(partes[1]), MESES.index(partes[0]) + 1, 1)


def generar(mes="auto", carpeta="."):
    """Valida la plantilla y el horario anterior, y crea el horario del mes en la carpeta salida/."""
    try:
        return _generar(mes, Path(carpeta))
    except ErrorUsuario as e:
        print(f"\n❌ No se pudo crear el horario.\n\n{e}")


def _generar(mes, base):
    data, salida = base / "data", base / "salida"
    plantilla = data / "PLANTILLA_MES.xlsx"
    inicio = None if mes == "auto" else _mes_desde_texto(mes)

    print("▶ Paso 1 de 4 · Revisando la plantilla (data/PLANTILLA_MES.xlsx)...")
    if not plantilla.exists():
        crear_plantilla(plantilla)
        raise ErrorUsuario("No existía la plantilla, así que se creó una nueva en data/PLANTILLA_MES.xlsx.\n"
                           "Ábrala, complete la hoja EJECUTIVOS, guarde, ciérrela y presione ▶ de nuevo.")
    ejecutivos, ausencias, reglas, feriados, avisos = leer_plantilla(plantilla)
    print(f"   ✅ {len(ejecutivos)} ejecutivos activos · {len(ausencias)} días de ausencia · turnos de {reglas['horas_semanales_contrato']:g} h correctos")

    print("▶ Paso 2 de 4 · Leyendo el horario anterior...")
    archivos = ([data / "horario_mes_anterior.xlsx"] if (data / "horario_mes_anterior.xlsx").exists() else []) \
        + sorted(salida.glob("PLANIFICACION_*.xlsx"), key=lambda p: p.stat().st_mtime)
    archivos = [a for a in archivos if not a.name.startswith("~$")]  # archivos temporales de Excel abierto
    historico = cargar_historico(archivos)
    if inicio is None:
        if historico.empty:
            raise ErrorUsuario("No se encontró ningún horario anterior para saber qué mes sigue.\n"
                               "Copie el Excel del mes pasado en la carpeta data/ con el nombre horario_mes_anterior.xlsx,\n"
                               "o escriba el mes en la casilla, por ejemplo MES = \"Diciembre 2026\".")
        inicio = historico["FECHA"].max().normalize() + pd.offsets.MonthBegin(1)
    historico = historico[historico["FECHA"] < inicio]
    nombre_mes = f"{MESES[inicio.month - 1].upper()} {inicio.year}"
    if historico.empty:
        avisos.append("No hay horario anterior a este mes: no se considera la rotación previa (sábados y tardes).")
        previo = {}
    else:
        previo = estado_previo(historico)
        print(f"   ✅ Horario anterior hasta el {historico['FECHA'].max():%d-%m-%Y} ({', '.join(a.name for a in archivos)})")
        nuevos = ejecutivos.loc[~ejecutivos["RUT"].isin(previo), "NOMBRE"].tolist()
        if nuevos:
            avisos.append(f"Sin historial (se tratan como nuevos): {', '.join(nuevos)}.")
    if not any(pd.Timestamp(f).year == inicio.year for f in feriados):
        avisos.append(f"La hoja FERIADOS no tiene feriados de {inicio.year}: agréguelos si corresponde.")
    print(f"   📅 Mes a planificar: {nombre_mes}")

    print("▶ Paso 3 de 4 · Asignando turnos, pausas y colaciones...")
    calendario = generar_calendario(inicio, feriados)
    turnos_semana = asignar_turnos_semanales(ejecutivos, calendario, previo, ausencias, reglas)
    plan = asignar_descansos(expandir_plan(ejecutivos, calendario, turnos_semana, ausencias, reglas), reglas)
    print(f"   ✅ {int((plan['ESTADO_DIA'] == 'TRABAJA').sum())} jornadas asignadas")

    print("▶ Paso 4 de 4 · Validando y creando el Excel...")
    alertas = auditar(plan, ejecutivos, previo, reglas)
    salida.mkdir(exist_ok=True)
    ruta = salida / f"PLANIFICACION_{nombre_mes.replace(' ', '_')}.xlsx"
    exportar_excel(ruta, plan, ejecutivos, turnos_semana, calendario, alertas, reglas)

    print(f"\n✅ LISTO. Horario de {nombre_mes} guardado en:\n   {ruta.resolve()}")
    if avisos:
        print("\nℹ️ Para su información:")
        for a in avisos:
            print(f"   • {a}")
    if alertas.empty:
        print("\n✅ El horario cumple todas las reglas.")
    else:
        print(f"\n⚠️ Hay {len(alertas)} alertas para revisar en la hoja ALERTAS del Excel:")
        for tipo, n in alertas["TIPO"].value_counts().items():
            print(f"   • {tipo}: {n}")
    return ruta
